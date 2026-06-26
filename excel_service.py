import io
from datetime import datetime, date
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


AVR_RED = 'C8102E'
AVR_BLACK = '1A1A1A'
AVR_GRAY = 'F2F2F2'


def export_customers_excel(customers):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Clienti'

    _write_header(ws, 'A.V.R. Informatica — Clienti', 10)

    headers = ['ID', 'Nome', 'Cognome', 'Azienda', 'Email', 'Telefono', 'Città', 'Abbonamenti Attivi', 'Creato il']
    _write_row_header(ws, 3, headers)

    for i, c in enumerate(customers, start=4):
        ws.cell(i, 1, c.id)
        ws.cell(i, 2, c.first_name)
        ws.cell(i, 3, c.last_name)
        ws.cell(i, 4, c.company or '')
        ws.cell(i, 5, c.email)
        ws.cell(i, 6, c.phone or '')
        ws.cell(i, 7, c.city or '')
        ws.cell(i, 8, c.active_subscriptions_count)
        ws.cell(i, 9, c.created_at.strftime('%d/%m/%Y'))
        if i % 2 == 0:
            _fill_row(ws, i, 9, AVR_GRAY)

    _auto_width(ws, 9)
    return _to_bytes(wb)


def export_subscriptions_excel(subscriptions):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Abbonamenti'

    _write_header(ws, 'A.V.R. Informatica — Abbonamenti', 14)

    headers = ['ID', 'Cliente', 'Azienda', 'Email', 'Licenza', 'Qtà', 'Inizio', 'Scadenza', 'Giorni Rimasti', 'Ciclo', 'Prezzo €', 'Stato']
    _write_row_header(ws, 3, headers)

    today = datetime.utcnow().date()
    for i, s in enumerate(subscriptions, start=4):
        days = (s.end_date - today).days
        ws.cell(i, 1, s.id)
        ws.cell(i, 2, s.customer.full_name)
        ws.cell(i, 3, s.customer.company or '')
        ws.cell(i, 4, s.customer.email)
        ws.cell(i, 5, s.license_type.name)
        ws.cell(i, 6, s.quantity)
        ws.cell(i, 7, s.start_date.strftime('%d/%m/%Y'))
        ws.cell(i, 8, s.end_date.strftime('%d/%m/%Y'))
        ws.cell(i, 9, days)
        ws.cell(i, 10, 'Mensile' if s.billing_cycle == 'monthly' else 'Annuale')
        ws.cell(i, 11, round(s.effective_price, 2))
        ws.cell(i, 12, s.status.capitalize())

        color = None
        if days < 0:
            color = 'FFCCCC'
        elif days <= 30:
            color = 'FFE0B2'
        elif days <= 90:
            color = 'FFF9C4'
        if color:
            _fill_row(ws, i, 12, color)
        elif i % 2 == 0:
            _fill_row(ws, i, 12, AVR_GRAY)

    _auto_width(ws, 12)
    return _to_bytes(wb)


def import_customers_excel(file_stream):
    df = pd.read_excel(file_stream, engine='openpyxl')
    df.columns = [str(c).strip().lower() for c in df.columns]

    COL_MAP = {
        'first_name': ['nome', 'first name', 'first_name', 'name'],
        'last_name': ['cognome', 'last name', 'last_name', 'surname'],
        'email': ['email', 'e-mail', 'mail'],
        'company': ['azienda', 'company', 'società', 'ragione sociale'],
        'phone': ['telefono', 'phone', 'tel', 'cellulare'],
        'city': ['città', 'city', 'comune'],
    }

    mapped = {}
    for field, variants in COL_MAP.items():
        for v in variants:
            if v in df.columns:
                mapped[field] = v
                break

    records = []
    errors = []
    for idx, row in df.iterrows():
        try:
            first = str(row.get(mapped.get('first_name', ''), '') or '').strip()
            last = str(row.get(mapped.get('last_name', ''), '') or '').strip()
            email = str(row.get(mapped.get('email', ''), '') or '').strip()
            if not first or not email:
                errors.append(f"Riga {idx+2}: nome o email mancante")
                continue
            records.append({
                'first_name': first,
                'last_name': last,
                'email': email,
                'company': str(row.get(mapped.get('company', ''), '') or '').strip(),
                'phone': str(row.get(mapped.get('phone', ''), '') or '').strip(),
                'city': str(row.get(mapped.get('city', ''), '') or '').strip(),
            })
        except Exception as e:
            errors.append(f"Riga {idx+2}: {e}")

    return records, errors


# ── helpers ──────────────────────────────────────────────────────────────────

def _write_header(ws, title, col_span):
    ws.merge_cells(f'A1:{get_column_letter(col_span)}1')
    cell = ws['A1']
    cell.value = title
    cell.font = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor=AVR_RED)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f'A2:{get_column_letter(col_span)}2')
    ws['A2'].value = f'Generato il {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    ws['A2'].font = Font(name='Calibri', italic=True, size=10, color='666666')
    ws['A2'].alignment = Alignment(horizontal='center')


def _write_row_header(ws, row, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row, col, h)
        c.font = Font(name='Calibri', bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AVR_BLACK)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin = Side(style='thin', color='FFFFFF')
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[row].height = 20


def _fill_row(ws, row, max_col, color):
    for col in range(1, max_col + 1):
        ws.cell(row, col).fill = PatternFill('solid', fgColor=color)


def _auto_width(ws, max_col):
    for col in range(1, max_col + 1):
        max_len = 0
        col_letter = get_column_letter(col)
        for cell in ws[col_letter]:
            try:
                max_len = max(max_len, len(str(cell.value or '')))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
